"""Success check (slide 17): same 24 h, keeper's sheet vs CivetWatch's sheet.

* sheets   - two observation sheets in the CivetWatch layout (a keeper can fill
             a blank copy of sheet.xlsx): hourly counts + Start/Stop durations
* events   - keeper events CSV (behavior,start[,stop]) vs events.csv
* frames   - frames.csv vs a per-sample ground truth (synthetic demo / labelled day)
"""
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from . import ethogram as E
from .postprocess import majority_vote

FREQ_COLS = ["Feeding", "Drinking", "Territorial marking", "Grooming", "Urine", "Feces"]
_ALIASES = {"marking": E.MARKING, "territorial marking": E.MARKING, "feces": E.DEFECATION,
            "defecation": E.DEFECATION, "urine": E.URINATION, "urination": E.URINATION}


def _norm_behavior(b):
    b = str(b).strip().lower()
    return _ALIASES.get(b, b.replace(" ", "_").replace("/", "_"))


def read_sheet(path):
    ws = load_workbook(path, data_only=True).worksheets[0]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    dur = next(k for k, v in cells.items() if v == "Duration")
    header = {}
    for (r, c), v in cells.items():
        if v in FREQ_COLS and r in (dur[0] + 1, dur[0] + 2):
            header[v] = c
    hourly, r = {}, dur[0] + 3
    while (r, dur[1]) in cells and cells[(r, dur[1])] != "Total":
        hourly[str(cells[(r, dur[1])])] = {k: int(cells.get((r, c)) or 0) for k, c in header.items()}
        r += 1
    blocks = {}
    for (r, c), v in cells.items():
        if r == 2 and v in ("Inactive", "Active", "Not visible"):
            times = []
            rr = 4
            while (rr, c) in cells:
                a, b = cells.get((rr, c + 1)), cells.get((rr, c + 2))
                if a and b:
                    times.append((str(a), str(b)))
                rr += 1
            blocks[v] = times
    return pd.DataFrame(hourly).T.reindex(columns=FREQ_COLS).fillna(0).astype(int), blocks


def _minutes(blocks):
    total = 0
    for a, b in blocks:
        ha, ma = [int(x) for x in a.split(":")[:2]]
        hb, mb = [int(x) for x in b.split(":")[:2]]
        total += ((hb * 60 + mb) - (ha * 60 + ma)) % (24 * 60)
    return total


def compare_sheets(keeper_path, auto_path):
    hk, bk = read_sheet(keeper_path)
    ha, ba = read_sheet(auto_path)
    hours = hk.index.union(ha.index, sort=False)
    hk, ha = hk.reindex(hours).fillna(0), ha.reindex(hours).fillna(0)
    rows = []
    for col in FREQ_COLS:
        k, a = int(hk[col].sum()), int(ha[col].sum())
        hourly_agree = float((hk[col] == ha[col]).mean()) if len(hours) else 1.0
        rows.append({"behavior": col, "keeper": k, "civetwatch": a, "difference": a - k,
                     "hours_matching": f"{hourly_agree:.0%}"})
    for state in ("Inactive", "Active", "Not visible"):
        k, a = _minutes(bk.get(state, [])), _minutes(ba.get(state, []))
        rows.append({"behavior": f"{state} (min)", "keeper": k, "civetwatch": a, "difference": a - k,
                     "hours_matching": ""})
    return pd.DataFrame(rows)


def compare_events(keeper_csv, events_csv, tolerance_s=60):
    keeper = pd.read_csv(keeper_csv)
    keeper["behavior"] = keeper["behavior"].map(_norm_behavior)
    keeper["start"] = pd.to_datetime(keeper["start"])
    auto = pd.read_csv(events_csv, parse_dates=["start", "stop"])
    auto = auto[auto["category"] == "frequency"]
    rows = []
    for beh in E.FREQUENCY_LABELS:
        k = keeper[keeper["behavior"] == beh].sort_values("start")
        a = auto[auto["behavior"] == beh].sort_values("start")
        used, hit = set(), 0
        for t in k["start"]:
            d = (a["start"] - t).abs().dt.total_seconds()
            for idx in d[d <= tolerance_s].sort_values().index:
                if idx not in used:
                    used.add(idx)
                    hit += 1
                    break
        rows.append({"behavior": E.DISPLAY[beh], "keeper": len(k), "civetwatch": len(a), "matched": hit,
                     "recall": hit / len(k) if len(k) else np.nan,
                     "precision": hit / len(a) if len(a) else np.nan})
    return pd.DataFrame(rows)


def compare_frames(frames_csv, truth_csv, vote_window=6):
    frames = pd.read_csv(frames_csv, parse_dates=["t"], keep_default_na=False)
    truth = pd.read_csv(truth_csv, parse_dates=["t"])
    if "voted" not in frames:
        frames["voted"] = majority_vote(frames["label"], frames["conf"], vote_window)
    m = pd.merge_asof(frames.sort_values("t"), truth.sort_values("t"), on="t",
                      tolerance=pd.Timedelta("300ms"), direction="nearest")
    m = m.dropna(subset=["true_label"])
    rows = []
    for lab in sorted(set(m["true_label"])):
        sel = m[m["true_label"] == lab]
        rows.append({"true label": E.DISPLAY.get(lab, lab), "samples": len(sel),
                     "correct": f"{(sel['voted'] == lab).mean():.0%}",
                     "most often called": E.DISPLAY.get(sel["voted"].mode().iloc[0], "")})
    st_true = m["true_label"].map(E.state_of)
    st_pred = m["state"] if "state" in m else m["voted"].map(E.state_of)
    known = st_true.notna()
    overall = {"label accuracy": float((m["voted"] == m["true_label"]).mean()),
               "state accuracy": float((st_true[known] == st_pred[known]).mean()),
               "room accuracy": float((m["room_x"] == m["room_y"]).mean()) if "room_y" in m else np.nan}
    return pd.DataFrame(rows), overall
