"""Output files: events.csv, frames.csv, sheet.xlsx (the zoo's observation
sheet layout), summary.json."""
import json
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import ethogram as E

THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="EDEDED")
DUR_FILL = PatternFill("solid", fgColor="E2EFDA")
FREQ_FILL = PatternFill("solid", fgColor="FFF2CC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
BOLD = Font(bold=True, size=9)
SMALL = Font(size=9)

EVENT_COLUMNS = ["start", "stop", "duration_s", "behavior", "category", "room", "camera", "confidence",
                 "light", "visibility", "quality", "note", "needs_review", "clip", "grid", "vlm_answer", "toilet_score"]


def write_csvs(result, out_dir):
    ev = result["events"].copy()
    if "clip" not in ev:
        ev["clip"] = ""
    ev = ev[[c for c in EVENT_COLUMNS if c in ev]]
    ev.insert(4, "behavior_name", ev["behavior"].map(lambda b: E.DISPLAY.get(b, b)))
    ev.to_csv(os.path.join(out_dir, "events.csv"), index=False, date_format="%Y-%m-%d %H:%M:%S")
    result["frames"].to_csv(os.path.join(out_dir, "frames.csv"), index=False,
                            date_format="%Y-%m-%d %H:%M:%S.%f")


def write_summary(result, out_dir, meta):
    s = dict(result["summary"])
    s["start"], s["end"] = s["start"].isoformat(), s["end"].isoformat()
    s["alerts"] = result["alerts"]
    s.update(meta)
    with open(os.path.join(out_dir, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(s, f, indent=2, ensure_ascii=False, default=str)


def _cell(ws, row, col, value=None, font=SMALL, fill=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.border, c.alignment, c.font = BORDER, CENTER, font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def _merge(ws, r1, c1, r2, c2, value, font=BOLD, fill=HEAD_FILL):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            _cell(ws, r, c, font=font, fill=fill)
    ws.cell(row=r1, column=c1, value=value)
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def write_sheet(result, path, cfg, cam_ids):
    blocks, events, hourly = result["blocks"], result["events"], result["hourly"]
    s = result["summary"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Observation sheet"

    day = s["start"]
    title = (f"ตารางบันทึกพฤติกรรมชะมดแผงสันหางดำ   ชื่อ: {cfg['civet_name']}   "
             f"วันที่: {day:%d/%m/%Y}   ผู้บันทึก: {cfg['recorder']}")
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=11)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=23)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # ---- left: Duration Start / Stop rows ----
    groups = [
        ("Inactive", blocks[blocks["state"] == E.INACTIVE] if len(blocks) else blocks),
        ("Active", blocks[blocks["state"] == E.ACTIVE] if len(blocks) else blocks),
        ("Mating Behavior", pd.DataFrame(columns=["start", "stop"])),   # one civet: always empty
        ("Not visible", events[events["behavior"] == E.NOT_VISIBLE] if len(events) else events),
    ]
    n_rows = max([8] + [len(g) for _, g in groups])
    for gi, (name, rows) in enumerate(groups):
        c0 = 1 + gi * 3
        _merge(ws, 2, c0, 2, c0 + 2, name)
        for j, h in enumerate(["No", "Start", "Stop"]):
            _cell(ws, 3, c0 + j, h, BOLD, HEAD_FILL)
        rows = rows.reset_index(drop=True)
        for i in range(n_rows):
            r = 4 + i
            _cell(ws, r, c0, i + 1)
            if i < len(rows):
                _cell(ws, r, c0 + 1, rows.loc[i, "start"].strftime("%H:%M"), fill=DUR_FILL)
                _cell(ws, r, c0 + 2, rows.loc[i, "stop"].strftime("%H:%M"), fill=DUR_FILL)
            else:
                _cell(ws, r, c0 + 1)
                _cell(ws, r, c0 + 2)

    # ---- right: hourly Frequency (Time) ----
    c0 = 14
    _merge(ws, 2, c0, 4, c0, "Duration")
    _merge(ws, 2, c0 + 1, 2, c0 + 7, "Frequency (Time)")
    for j, h in enumerate(["Feeding", "Drinking", "Territorial marking", "Grooming"]):
        _merge(ws, 3, c0 + 1 + j, 4, c0 + 1 + j, h)
    _merge(ws, 3, c0 + 5, 3, c0 + 6, "Excretion")
    for j, h in enumerate(["Urine", "Feces"]):
        _cell(ws, 4, c0 + 5 + j, h, BOLD, HEAD_FILL)
    _merge(ws, 3, c0 + 7, 4, c0 + 7, "Total")
    cols = ["Feeding", "Drinking", "Territorial marking", "Grooming", "Urine", "Feces", "Total"]
    r = 5
    for _, row in hourly.iterrows():
        _cell(ws, r, c0, row["hour"], BOLD)
        for j, col in enumerate(cols):
            v = int(row[col])
            fill = FREQ_FILL if v and col != "Total" else None
            _cell(ws, r, c0 + 1 + j, v if v or col == "Total" else None, BOLD if col == "Total" else SMALL, fill)
        r += 1
    _cell(ws, r, c0, "Total", BOLD, HEAD_FILL)
    for j, col in enumerate(cols):
        _cell(ws, r, c0 + 1 + j, int(hourly[col].sum()) if len(hourly) else 0, BOLD, HEAD_FILL)

    note_row = max(r, 4 + n_rows) + 2
    cams = ", ".join(cam_ids)
    ws.cell(row=note_row, column=1, value=(
        f"Filled automatically by CivetWatch from {cams} · {cfg['sampling']['fps']} frames/second · "
        f"Inactive/Active switch after ≥{cfg['postprocess']['min_state_s'] // 60:g} min · "
        "meaning of every word: 'Read me' tab")).font = Font(italic=True, size=8)

    for c in range(1, 13):
        ws.column_dimensions[get_column_letter(c)].width = 6.5
    ws.column_dimensions["M"].width = 2
    ws.column_dimensions[get_column_letter(c0)].width = 12
    for c in range(c0 + 1, c0 + 9):
        ws.column_dimensions[get_column_letter(c)].width = 10
    ws.row_dimensions[3].height = 28

    # ---- other sheets ----
    _table_sheet(wb, "Events", _events_for_sheet(events))
    rest = events[events["behavior"].isin(E.INACTIVE_LABELS)] if len(events) else events
    _table_sheet(wb, "Inactive detail", _events_for_sheet(rest))
    stops = events[events["category"] == "stop"] if len(events) else events
    if len(stops):
        _table_sheet(wb, "Stops to check", _events_for_sheet(stops.sort_values("needs_review", ascending=False)))
    _table_sheet(wb, "Summary", pd.DataFrame(_summary_rows(s), columns=["Item", "Value"]))
    _table_sheet(wb, "Alerts", pd.DataFrame({"Alert": result["alerts"] or ["(none)"]}))
    from .glossary import GLOSSARY
    _table_sheet(wb, "Read me", pd.DataFrame(GLOSSARY, columns=["Word", "What it means"]))
    rm = wb["Read me"]
    rm.column_dimensions["A"].width, rm.column_dimensions["B"].width = 34, 120
    for row in rm.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    wb.save(path)


def _events_for_sheet(events):
    if not len(events):
        return pd.DataFrame(columns=["Start", "Stop", "Behavior"])
    out = pd.DataFrame({
        "Start": events["start"].dt.strftime("%Y-%m-%d %H:%M:%S"),
        "Stop": events["stop"].dt.strftime("%H:%M:%S"),
        "Minutes": (events["duration_s"] / 60).round(2),
        "Behavior": events["behavior"].map(lambda b: E.DISPLAY.get(b, b)),
        "Counted as": events["category"],
        "Room": events["room"], "Camera": events["camera"],
        "Confidence": events["confidence"],
        "Condition": events["light"] + ", " + events["visibility"] + ", " + events["quality"],
        "Note": events.get("note", ""),
        "Check": events["needs_review"].map({True: "please check", False: ""}),
    })
    if "clip" in events:
        out["Clip"] = events["clip"].fillna("")
    return out


def fmt_hm(seconds):
    m = int(round(seconds / 60))
    return f"{m // 60} h {m % 60:02d} m"


def _summary_rows(s):
    rows = [
        ("Observed", f"{s['start']:%Y-%m-%d %H:%M} → {s['end']:%Y-%m-%d %H:%M} ({fmt_hm(s['observed_s'])})"),
        ("Active", f"{fmt_hm(s['active_s'])} ({s['active_frac']:.0%})"),
        ("Inactive", fmt_hm(s["inactive_s"])),
        ("Not visible", fmt_hm(s["not_visible_s"])),
        ("Not seen by any camera (state unknown)", fmt_hm(s.get("unseen_s", 0))),
        ("Night (IR) frames", f"{s['night_frac']:.0%}"),
        ("Poor-image frames", f"{s['poor_image_frac']:.1%}"),
        ("Events to review", s["to_review"]),
        ("Short stops (possible toileting / marking / grooming)", s.get("short_stops", 0)),
        ("  flagged by the model for checking", s.get("stops_flagged", 0)),
    ]
    rows += [(f"Time in {room}", f"{share:.0%}") for room, share in s["room_share"].items()]
    rows += [(E.DISPLAY.get(k, k), fmt_hm(v)) for k, v in s["rest_types_s"].items()]
    rows += [(f"{E.DISPLAY.get(k, k)} (count)", v) for k, v in s["counts"].items()]
    return rows


def _table_sheet(wb, name, df):
    ws = wb.create_sheet(name)
    for j, col in enumerate(df.columns, 1):
        _cell(ws, 1, j, col, BOLD, HEAD_FILL)
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(col)) + 4))
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
            c.font = SMALL
    if name in ("Events", "Inactive detail"):
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["I"].width = 30
    ws.freeze_panes = "A2"
